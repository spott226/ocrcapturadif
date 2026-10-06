"""Importa una exportación de Captura Apoyos sin mostrar datos personales."""
from __future__ import annotations

import argparse
import os
import re
import sys
import unicodedata
from datetime import datetime, timezone
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
os.chdir(ROOT)
sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook
from sqlalchemy import inspect, or_, select

from app.database import SessionLocal, engine
from app.models import Person, SupportType


def normalized_header(value) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(char for char in text if not unicodedata.combining(char))
    return re.sub(r"[^A-Z0-9]+", " ", text.upper()).strip()


HEADERS = {
    "ID": "source_id",
    "NOMBRE COMPLETO": "name",
    "NOMBRE S": "given_names",
    "NOMBRES": "given_names",
    "APELLIDO PATERNO": "paternal_surname",
    "APELLIDO MATERNO": "maternal_surname",
    "CURP": "curp",
    "DIRECCION": "address",
    "MUNICIPIO": "municipality",
    "CODIGO DE APOYO": "support_code",
    "CODIGO TIPO DE APOYO": "support_code",
    "TIPO DE APOYO": "support_name",
    "NUMERO DE TELEFONO": "phone",
    "TELEFONO": "phone",
    "LIDER": "leader",
    "CAPTURO": "created_by",
    "FECHA DE CAPTURA": "created_at",
}


def as_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def parse_date(value) -> datetime:
    if isinstance(value, datetime):
        return value if value.tzinfo else value.replace(tzinfo=timezone.utc)
    text = as_text(value)
    for pattern in ("%d/%m/%Y %H:%M", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, pattern).replace(tzinfo=timezone.utc)
        except ValueError:
            pass
    return datetime.now(timezone.utc)


def clean_upper(value, limit: int) -> str:
    return " ".join(as_text(value).upper().split())[:limit]


def normalized_value(value) -> str:
    return normalized_header(as_text(value))


def import_workbook(path: Path, simulate: bool = False) -> tuple[int, int, int]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    try:
        raw_headers = next(rows)
    except StopIteration as exc:
        raise ValueError("El Excel está vacío") from exc

    columns = {index: HEADERS.get(normalized_header(value)) for index, value in enumerate(raw_headers)}
    available_fields = set(columns.values())
    is_v2 = "given_names" in available_fields or "paternal_surname" in available_fields
    if "name" not in available_fields and not {
        "given_names", "paternal_surname"
    }.issubset(available_fields):
        raise ValueError(
            "No se encontró 'Nombre completo' ni las columnas Nombre(s) y Apellido paterno"
        )

    required_tables = {"people", "support_types"}
    missing_tables = required_tables - set(inspect(engine).get_table_names())
    if missing_tables:
        raise RuntimeError(
            "Falta aplicar la migración SQL v2; tablas ausentes: "
            + ", ".join(sorted(missing_tables))
        )

    inserted = duplicates = invalid = 0
    seen_curps: set[str] = set()
    seen_phones: set[str] = set()
    with SessionLocal() as db:
        support_items = list(db.scalars(select(SupportType)).all())
        supports_by_code = {
            normalized_value(item.code): item for item in support_items
        }
        supports_by_name: dict[str, SupportType | None] = {}
        for item in support_items:
            key = normalized_value(item.name)
            supports_by_name[key] = item if key not in supports_by_name else None

        for row in rows:
            values = {
                field: row[index]
                for index, field in columns.items()
                if field and index < len(row)
            }
            given_names = clean_upper(values.get("given_names"), 120)
            paternal_surname = clean_upper(values.get("paternal_surname"), 80)
            maternal_surname = clean_upper(values.get("maternal_surname"), 80)
            legacy_name = clean_upper(values.get("name"), 180)
            if is_v2:
                name = " ".join(
                    part for part in (given_names, paternal_surname, maternal_surname) if part
                )[:180]
            else:
                name = legacy_name
            curp = clean_upper(values.get("curp"), 18)
            phone = re.sub(r"\D", "", as_text(values.get("phone")))[:15]
            municipality = " ".join(as_text(values.get("municipality")).split())[:120]

            support = None
            support_reference_present = bool(
                as_text(values.get("support_code")) or as_text(values.get("support_name"))
            )
            if as_text(values.get("support_code")):
                support = supports_by_code.get(normalized_value(values.get("support_code")))
            elif as_text(values.get("support_name")):
                support = supports_by_name.get(normalized_value(values.get("support_name")))

            if (
                not name
                or (is_v2 and (not given_names or not paternal_surname))
                or (support_reference_present and support is None)
            ):
                invalid += 1
                continue

            if (curp and curp in seen_curps) or (phone and phone in seen_phones):
                duplicates += 1
                continue

            conditions = []
            if curp:
                conditions.append(Person.curp == curp)
            if phone:
                conditions.append(Person.phone == phone)
            if conditions and db.scalar(select(Person.id).where(or_(*conditions)).limit(1)):
                duplicates += 1
                continue

            if not simulate:
                db.add(Person(
                    name=name,
                    given_names=given_names,
                    paternal_surname=paternal_surname,
                    maternal_surname=maternal_surname,
                    curp=curp,
                    address=("" if is_v2 else clean_upper(values.get("address"), 500)),
                    municipality=municipality,
                    support_type_id=support.id if support else None,
                    phone=phone,
                    leader=clean_upper(values.get("leader"), 180),
                    created_by=as_text(values.get("created_by"))[:254] or "IMPORTACION",
                    created_at=parse_date(values.get("created_at")),
                ))
            if curp:
                seen_curps.add(curp)
            if phone:
                seen_phones.add(phone)
            inserted += 1

        if not simulate:
            db.commit()
    workbook.close()
    return inserted, duplicates, invalid


def main() -> int:
    parser = argparse.ArgumentParser(description="Importa una exportación .xlsx de Captura Apoyos")
    parser.add_argument("archivo", type=Path)
    parser.add_argument("--simular", action="store_true", help="valida todo sin guardar")
    args = parser.parse_args()
    if not args.archivo.is_file() or args.archivo.suffix.casefold() != ".xlsx":
        parser.error("indique un archivo .xlsx existente")
    inserted, duplicates, invalid = import_workbook(args.archivo, args.simular)
    mode = "SIMULACION" if args.simular else "IMPORTACION"
    print(f"{mode} TERMINADA: nuevos={inserted}, duplicados={duplicates}, inválidos={invalid}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
