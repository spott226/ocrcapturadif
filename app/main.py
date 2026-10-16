import logging
import re
import unicodedata
from contextlib import asynccontextmanager
from datetime import date, datetime, timezone
from io import BytesIO
from pathlib import Path

from fastapi import Depends, FastAPI, File, Form, HTTPException, Query, Request, UploadFile
from fastapi.responses import HTMLResponse, JSONResponse, RedirectResponse, Response, StreamingResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from openpyxl import Workbook
from sqlalchemy import DateTime, Integer, Unicode, func, inspect, or_, select, text, true
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session
from starlette.middleware.httpsredirect import HTTPSRedirectMiddleware
from starlette.middleware.sessions import SessionMiddleware

from .config import get_settings
from .database import Base, SessionLocal, engine, get_db
from .models import AppUser, Person, StructureSubmission, SupportNetworkSubmission, SupportType
from .ocr import extract_image
from .security import csrf_token, password_hash, valid_csrf, verify_password

settings = get_settings()
BASE = Path(__file__).resolve().parent
logger = logging.getLogger("uvicorn.error")


def normalize_public_prefix(value: str) -> str:
    prefix = value.strip()
    if not prefix or prefix == "/":
        return ""
    return "/" + prefix.strip("/")


PUBLIC_PREFIX = normalize_public_prefix(settings.public_prefix)


def public_url(path: str = "/") -> str:
    if not path.startswith("/"):
        path = "/" + path
    return f"{PUBLIC_PREFIX}{path}"

ROLE_SUPERADMIN = "superadmin"
ROLE_CAPTURISTA = "capturista"
VALID_ROLES = {ROLE_SUPERADMIN, ROLE_CAPTURISTA}

FORM_OCR = "ocr"
FORM_GOOGLE_1 = "google_1"
FORM_RED_1X10 = "red_1x10"
FORM_DEFINITIONS = (
    {
        "code": FORM_OCR,
        "name": "Captura de INE con OCR",
        "description": "Fotografíe la INE, revise los datos detectados y guarde la captura.",
        "url": "/captura",
        "external": False,
    },
    {
        "code": FORM_GOOGLE_1,
        "name": "Datos de estructura",
        "description": "Información relevante sobre los integrantes del equipo, con frente y reverso de INE.",
        "url": "/formularios/datos-estructura",
        "external": False,
    },
    {
        "code": FORM_RED_1X10,
        "name": "Red De Apoyo 1x10",
        "description": "Formulario de registro de la red de apoyo.",
        "url": "/formularios/red-apoyo",
        "external": False,
    },
)
VALID_FORM_CODES = {item["code"] for item in FORM_DEFINITIONS}

NETWORK_LEADERS = (
    "Arath Gomez", "Yeraldine Baez", "Alejandro Santillan", "Arturo Buenrostro",
    "Bladimir Hernandez", "Cesar", "Christopher Sanchez", "Jazmin",
    "Jimchauregy Baez", "Montserrat (fotografa)", "Mario Gutierrez", "Nayeli Luevano",
    "Olivia Gonzalez", "Paloma", "Pedro", "Raudel Calvillo",
    "Seleste Montserrat (Hacienda)", "Yordin Cervantes", "Salma Ovalle",
    "Chico de Luna", "Derek Legaspy", "Ilse", "Fabian", "Gael Luevano",
)

MUNICIPALITIES = (
    "Aguascalientes",
    "Asientos",
    "Calvillo",
    "Cosío",
    "El Llano",
    "Jesús María",
    "Pabellón de Arteaga",
    "Rincón de Romos",
    "San Francisco de los Romo",
    "San José de Gracia",
    "Tepezalá",
)

LEADERS = (
    "Nalleli Luévano",
    "Jose Guadalupe Barajas Luevanos",
    "Seleste Monserrat Contreras Rodriguez",
    "Ma. Olivia Mercedes González Rojas",
    "Jimchauregi Gutiérrez Baez",
    "Shelsie Yeraldine López Báez",
    "Derek Legaspy Reyes Cervantes",
    "Jorge Fabian Luévano Rodríguez",
    "Brenda Nallely Ortiz Herrera",
    "Nanci Nalleli Luévano Rico",
    "Arturo Alejandro Álvarez Buenrostro",
    "Gael Oziel Luévano Calvillo",
    "Cristopher Yahvé Moreno Sánchez",
    "Salma Dayana Ovalle Aguilar",
    "Diana Montserrat Luevano Puentes",
    "José Gustavo Díaz Contreras",
)

# These additions keep a database created by a previous release readable. The
# deployment package also carries an audited SQL Server migration; this small
# migration remains useful for local SQLite development and safe restarts.
EXTRA_COLUMNS = {
    "phone": Unicode(15),
    "leader": Unicode(180),
    "birth_date": Unicode(20),
    "sex_or_gender": Unicode(20),
    "state_code": Unicode(20),
    "municipality_code": Unicode(20),
    "section": Unicode(10),
    "locality_code": Unicode(20),
    "registration_year": Unicode(20),
    "issue_year": Unicode(10),
    "cic": Unicode(20),
    "ocr_code": Unicode(20),
    "given_names": Unicode(120),
    "paternal_surname": Unicode(80),
    "maternal_surname": Unicode(80),
    "municipality": Unicode(120),
}


def migrate_existing_database() -> None:
    inspector = inspect(engine)
    if "people" not in inspector.get_table_names():
        return
    existing = {column["name"] for column in inspector.get_columns("people")}
    with engine.begin() as connection:
        add_keyword = "ADD" if engine.dialect.name == "mssql" else "ADD COLUMN"
        for column, column_type in EXTRA_COLUMNS.items():
            if column not in existing:
                type_sql = column_type.compile(dialect=engine.dialect)
                connection.execute(text(
                    f"ALTER TABLE people {add_keyword} {column} {type_sql} NOT NULL DEFAULT ''"
                ))
        if "support_type_id" not in existing:
            type_sql = Integer().compile(dialect=engine.dialect)
            connection.execute(text(
                f"ALTER TABLE people {add_keyword} support_type_id {type_sql} NULL"
            ))


def migrate_early_v2_database() -> None:
    """Keep local/preview databases made by an early v2 build usable."""
    inspector = inspect(engine)
    tables = set(inspector.get_table_names())
    with engine.begin() as connection:
        add_keyword = "ADD" if engine.dialect.name == "mssql" else "ADD COLUMN"
        for table_name in ("support_types", "app_users"):
            if table_name not in tables:
                continue
            columns = {column["name"] for column in inspector.get_columns(table_name)}
            if "active" in columns and "is_active" not in columns:
                if engine.dialect.name == "mssql":
                    connection.execute(text(
                        f"EXEC sp_rename N'dbo.{table_name}.active', N'is_active', N'COLUMN'"
                    ))
                else:
                    connection.execute(text(
                        f"ALTER TABLE {table_name} RENAME COLUMN active TO is_active"
                    ))
        if "app_users" in tables:
            columns = {column["name"] for column in inspect(engine).get_columns("app_users")}
            if "last_login_at" not in columns:
                type_sql = DateTime(timezone=True).compile(dialect=engine.dialect)
                connection.execute(text(
                    f"ALTER TABLE app_users {add_keyword} last_login_at {type_sql} NULL"
                ))
            if "form_access" not in columns:
                type_sql = Unicode(255).compile(dialect=engine.dialect)
                default_sql = "N'ocr'" if engine.dialect.name == "mssql" else "'ocr'"
                connection.execute(text(
                    f"ALTER TABLE app_users {add_keyword} form_access {type_sql} "
                    f"NOT NULL DEFAULT {default_sql}"
                ))
        if "support_network_submissions" in tables:
            columns = {column["name"] for column in inspect(engine).get_columns("support_network_submissions")}
            type_sql = Unicode(300).compile(dialect=engine.dialect)
            default_sql = "N''" if engine.dialect.name == "mssql" else "''"
            for number in range(1, 11):
                column = f"facebook_{number}"
                if column not in columns:
                    connection.execute(text(
                        f"ALTER TABLE support_network_submissions {add_keyword} {column} {type_sql} "
                        f"NOT NULL DEFAULT {default_sql}"
                    ))


def bootstrap_first_superadmin() -> None:
    """Import the existing environment administrator once, never on every boot."""
    with SessionLocal() as db:
        if (db.scalar(select(func.count(AppUser.id))) or 0) != 0:
            return
        db.add(AppUser(
            email=settings.admin_email.strip().casefold(),
            password_hash=password_hash.hash(settings.admin_password),
            role=ROLE_SUPERADMIN,
            active=True,
            created_by="bootstrap",
        ))
        try:
            db.commit()
        except IntegrityError:
            # A second worker may have completed the same one-time bootstrap.
            db.rollback()


@asynccontextmanager
async def lifespan(_app: FastAPI):
    Base.metadata.create_all(engine)
    migrate_existing_database()
    migrate_early_v2_database()
    bootstrap_first_superadmin()
    yield


app = FastAPI(
    title="Captura Apoyos",
    docs_url=None,
    redoc_url=None,
    # IIS removes the application path before forwarding to HttpPlatformHandler.
    # PUBLIC_PREFIX is therefore only for browser-facing links and redirects.
    # Setting it as ASGI root_path makes Starlette strip the prefix a second time
    # and causes mounted static files to return 404.
    root_path="",
    lifespan=lifespan,
)
app.add_middleware(
    SessionMiddleware,
    secret_key=settings.secret_key,
    https_only=settings.cookie_secure,
    same_site="lax",
    max_age=28800,
)
if settings.force_https:
    app.add_middleware(HTTPSRedirectMiddleware)
app.mount("/static", StaticFiles(directory=BASE / "static"), name="static")
templates = Jinja2Templates(directory=BASE / "templates")


@app.exception_handler(HTTPException)
async def friendly_http_errors(request: Request, exc: HTTPException):
    if exc.status_code == 401:
        request.session.clear()
        return RedirectResponse(public_url("/login"), status_code=303)
    return JSONResponse({"detail": exc.detail}, status_code=exc.status_code, headers=exc.headers)


@app.middleware("http")
async def privacy_headers(request: Request, call_next):
    response = await call_next(request)
    response.headers["X-Content-Type-Options"] = "nosniff"
    response.headers["X-Frame-Options"] = "DENY"
    response.headers["Referrer-Policy"] = "no-referrer"
    response.headers["Permissions-Policy"] = "camera=(self), geolocation=(), microphone=()"
    response.headers["Content-Security-Policy"] = (
        "default-src 'self'; img-src 'self' data: blob:; media-src 'self' blob:; "
        "style-src 'self'; script-src 'self'; form-action 'self'; frame-ancestors 'none'"
    )
    if request.url.scheme == "https":
        response.headers["Strict-Transport-Security"] = "max-age=31536000; includeSubDomains"
    if request.url.path != "/salud" and not request.url.path.startswith("/static/"):
        response.headers["Cache-Control"] = "no-store"
    return response


def page(request: Request, name: str, **context):
    user = getattr(request.state, "current_user", None)
    context.update(
        request=request,
        user=user,
        csrf=csrf_token(request.session),
        public_prefix=PUBLIC_PREFIX,
    )
    if user and "can_use_ocr" not in context:
        context["can_use_ocr"] = FORM_OCR in allowed_form_codes(user)
    return templates.TemplateResponse(request, name, context)


def current_user(request: Request, db: Session) -> AppUser | None:
    """Reload the user on every request so role/active changes apply immediately."""
    raw_user_id = request.session.get("user_id")
    try:
        user_id = int(raw_user_id)
    except (TypeError, ValueError):
        return None
    user = db.get(AppUser, user_id)
    if not user or not user.active:
        request.session.clear()
        return None
    request.state.current_user = user
    return user


def require_user(request: Request, db: Session) -> AppUser:
    user = current_user(request, db)
    if not user:
        raise HTTPException(401, "Inicie sesión")
    return user


def require_superadmin(request: Request, db: Session) -> AppUser:
    user = require_user(request, db)
    if not user.is_superadmin:
        raise HTTPException(403, "Esta acción requiere el rol de superadministrador")
    return user


def normalized_form_codes(values: list[str] | tuple[str, ...] | None) -> list[str]:
    requested = set(values or [])
    return [item["code"] for item in FORM_DEFINITIONS if item["code"] in requested]


def allowed_form_codes(user: AppUser) -> set[str]:
    if user.is_superadmin:
        return set(VALID_FORM_CODES)
    # Existing capturistas created before this feature keep their OCR access.
    return user.form_codes or {FORM_OCR}


def forms_for_user(user: AppUser) -> list[dict]:
    allowed = allowed_form_codes(user)
    return [item for item in FORM_DEFINITIONS if item["code"] in allowed]


def require_form_access(request: Request, db: Session, form_code: str) -> AppUser:
    user = require_user(request, db)
    if form_code not in allowed_form_codes(user):
        raise HTTPException(403, "Su cuenta no tiene acceso a este formulario")
    return user


def require_csrf(request: Request, token: str) -> None:
    if not valid_csrf(request.session, token):
        raise HTTPException(403, "Solicitud no válida")


def active_support_types(db: Session) -> list[SupportType]:
    return list(db.scalars(
        select(SupportType)
        .where(SupportType.active == true())
        .order_by(SupportType.sort_order, SupportType.name)
    ).all())


def normalize_for_match(value: str) -> str:
    value = unicodedata.normalize("NFKD", value.upper())
    return " ".join("".join(char for char in value if not unicodedata.combining(char)).split())


def infer_municipality(address: str) -> str:
    normalized_address = normalize_for_match(address)
    for municipality in sorted(MUNICIPALITIES, key=len, reverse=True):
        if normalize_for_match(municipality) in normalized_address:
            return municipality
    return ""


def canonical_municipality(value: str) -> str:
    normalized = normalize_for_match(value)
    return next(
        (item for item in MUNICIPALITIES if normalize_for_match(item) == normalized),
        "",
    )


def canonical_leader(value: str) -> str:
    normalized = normalize_for_match(value)
    return next(
        (item for item in LEADERS if normalize_for_match(item) == normalized),
        "",
    )


def split_ine_name(full_name: str) -> tuple[str, str, str]:
    """Use the INE front order: paternal, maternal, then one or more given names."""
    words = full_name.strip().upper().split()
    if len(words) >= 3:
        return " ".join(words[2:]), words[0], words[1]
    if len(words) == 2:
        return words[1], words[0], ""
    if words:
        return words[0], "", ""
    return "", "", ""


def review_page(
    request: Request,
    db: Session,
    data: dict,
    *,
    error: str | None = None,
    duplicates: list[Person] | None = None,
):
    return page(
        request,
        "review.html",
        data=data,
        error=error,
        duplicates=duplicates or [],
        municipalities=MUNICIPALITIES,
        leaders=LEADERS,
        support_types=active_support_types(db),
    )


@app.get("/salud")
def health():
    return {"status": "ok"}


@app.get("/login", response_class=HTMLResponse)
def login_form(
    request: Request,
    cambiada: int = Query(0),
    db: Session = Depends(get_db),
):
    if current_user(request, db):
        return RedirectResponse(public_url("/"), status_code=303)
    return page(
        request,
        "login.html",
        success="Contraseña actualizada. Ingrese con la nueva contraseña." if cambiada else None,
    )


@app.post("/login", response_class=HTMLResponse)
def login(
    request: Request,
    email: str = Form(...),
    password: str = Form(...),
    csrf: str = Form(...),
    db: Session = Depends(get_db),
):
    require_csrf(request, csrf)
    normalized_email = email.strip().casefold()
    user = db.scalar(select(AppUser).where(func.lower(AppUser.email) == normalized_email))
    if not user or not user.active or not verify_password(password, user.password_hash):
        return page(request, "login.html", error="Correo o contraseña incorrectos")
    user.last_login_at = datetime.now(timezone.utc)
    db.commit()
    request.session.clear()
    request.session["user_id"] = user.id
    return RedirectResponse(public_url("/"), status_code=303)


@app.post("/logout")
def logout(request: Request, csrf: str = Form(...)):
    require_csrf(request, csrf)
    request.session.clear()
    return RedirectResponse(public_url("/login"), status_code=303)


def password_page(request: Request, *, error: str | None = None):
    return page(request, "password.html", error=error)


@app.get("/cuenta/contrasena", response_class=HTMLResponse)
def change_password_form(request: Request, db: Session = Depends(get_db)):
    require_user(request, db)
    return password_page(request)


@app.post("/cuenta/contrasena", response_class=HTMLResponse)
def change_password(
    request: Request,
    current_password: str = Form(...),
    new_password: str = Form(...),
    confirm_password: str = Form(...),
    csrf: str = Form(...),
    db: Session = Depends(get_db),
):
    user = require_user(request, db)
    require_csrf(request, csrf)
    if not verify_password(current_password, user.password_hash):
        return password_page(request, error="La contraseña actual no es correcta.")
    if not 12 <= len(new_password) <= 200:
        return password_page(
            request,
            error="La nueva contraseña debe tener entre 12 y 200 caracteres.",
        )
    if new_password != confirm_password:
        return password_page(request, error="La confirmación no coincide.")
    if verify_password(new_password, user.password_hash):
        return password_page(request, error="La nueva contraseña debe ser diferente.")
    user.password_hash = password_hash.hash(new_password)
    db.commit()
    request.session.clear()
    return RedirectResponse(public_url("/login?cambiada=1"), status_code=303)


@app.get("/", response_class=HTMLResponse)
def home(request: Request, db: Session = Depends(get_db)):
    user = require_user(request, db)
    return page(request, "index.html", available_forms=forms_for_user(user))


@app.get("/captura", response_class=HTMLResponse)
def capture_page(request: Request, db: Session = Depends(get_db)):
    user = require_form_access(request, db, FORM_OCR)
    return page(request, "capture.html", user=user, support_types=active_support_types(db))


@app.get("/beneficiarios", response_class=HTMLResponse)
def beneficiaries_page(
    request: Request,
    pagina: int = Query(1, ge=1),
    db: Session = Depends(get_db),
):
    require_superadmin(request, db)
    page_size = 10
    total = db.scalar(select(func.count(Person.id))) or 0
    total_pages = max(1, (total + page_size - 1) // page_size)
    current_page = min(pagina, total_pages)
    people = db.scalars(
        select(Person)
        .order_by(Person.created_at.desc())
        .offset((current_page - 1) * page_size)
        .limit(page_size)
    ).unique().all()
    return page(
        request,
        "beneficiaries.html",
        people=people,
        total=total,
        current_page=current_page,
        total_pages=total_pages,
    )


async def read_upload(upload: UploadFile) -> bytes:
    allowed = {"image/jpeg", "image/png", "image/webp"}
    if upload.content_type not in allowed:
        raise HTTPException(400, "Use una imagen JPG, PNG o WEBP")
    data = await upload.read(settings.max_upload_mb * 1024 * 1024 + 1)
    await upload.close()
    if len(data) > settings.max_upload_mb * 1024 * 1024:
        raise HTTPException(413, "La imagen excede el tamaño permitido")
    return data


@app.get("/formularios/datos-estructura", response_class=HTMLResponse)
def structure_form(request: Request, db: Session = Depends(get_db)):
    require_form_access(request, db, FORM_GOOGLE_1)
    return page(request, "structure_form.html")


@app.post("/formularios/datos-estructura", response_class=HTMLResponse)
async def save_structure_form(
    request: Request,
    email: str = Form(...),
    full_name: str = Form(...),
    birth_date: date = Form(...),
    electoral_section: str = Form(...),
    phone: str = Form(...),
    facebook_profile: str = Form(...),
    education: str = Form(...),
    age: int = Form(...),
    workplace: str = Form(...),
    front: UploadFile = File(...),
    back: UploadFile = File(...),
    csrf: str = Form(...),
    db: Session = Depends(get_db),
):
    user = require_form_access(request, db, FORM_GOOGLE_1)
    require_csrf(request, csrf)
    normalized_email = email.strip().casefold()
    normalized_phone = re.sub(r"\D", "", phone)
    normalized_section = re.sub(r"\D", "", electoral_section)
    if not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", normalized_email):
        raise HTTPException(400, "Capture un correo válido")
    if not full_name.strip() or not education.strip() or not workplace.strip():
        raise HTTPException(400, "Complete todos los campos obligatorios")
    if len(normalized_phone) != 10:
        raise HTTPException(400, "El número telefónico debe tener 10 dígitos")
    if not normalized_section or len(normalized_section) > 10:
        raise HTTPException(400, "Capture una sección electoral válida")
    if age < 15 or age > 120:
        raise HTTPException(400, "Capture una edad válida")
    front_bytes = await read_upload(front)
    back_bytes = await read_upload(back)
    db.add(StructureSubmission(
        email=normalized_email,
        full_name=full_name.strip()[:220],
        birth_date=birth_date,
        electoral_section=normalized_section,
        phone=normalized_phone,
        facebook_profile=facebook_profile.strip()[:300],
        education=education.strip()[:220],
        age=age,
        workplace=workplace.strip()[:300],
        front_image=front_bytes,
        front_mime=front.content_type or "image/jpeg",
        back_image=back_bytes,
        back_mime=back.content_type or "image/jpeg",
        created_by=user.email,
    ))
    db.commit()
    return RedirectResponse(public_url("/formularios/datos-estructura?guardado=1"), status_code=303)


@app.get("/formularios/red-apoyo", response_class=HTMLResponse)
def support_network_form(request: Request, db: Session = Depends(get_db)):
    require_form_access(request, db, FORM_RED_1X10)
    return page(request, "support_network_form.html", network_leaders=NETWORK_LEADERS)


@app.post("/formularios/red-apoyo", response_class=HTMLResponse)
def save_support_network_form(
    request: Request,
    leadership: str = Form(...),
    person_1: str = Form(...), person_2: str = Form(...),
    person_3: str = Form(...), person_4: str = Form(...),
    person_5: str = Form(...), person_6: str = Form(...),
    person_7: str = Form(...), person_8: str = Form(...),
    person_9: str = Form(...), person_10: str = Form(...),
    facebook_1: str = Form(...), facebook_2: str = Form(...),
    facebook_3: str = Form(...), facebook_4: str = Form(...),
    facebook_5: str = Form(...), facebook_6: str = Form(...),
    facebook_7: str = Form(...), facebook_8: str = Form(...),
    facebook_9: str = Form(...), facebook_10: str = Form(...),
    csrf: str = Form(...),
    db: Session = Depends(get_db),
):
    user = require_form_access(request, db, FORM_RED_1X10)
    require_csrf(request, csrf)
    if leadership not in NETWORK_LEADERS:
        raise HTTPException(400, "Seleccione un liderazgo válido")
    people = [person_1, person_2, person_3, person_4, person_5,
              person_6, person_7, person_8, person_9, person_10]
    facebook_profiles = [facebook_1, facebook_2, facebook_3, facebook_4, facebook_5,
                         facebook_6, facebook_7, facebook_8, facebook_9, facebook_10]
    if any(not value.strip() for value in people + facebook_profiles):
        raise HTTPException(400, "Capture el nombre y Facebook de las diez personas")
    db.add(SupportNetworkSubmission(
        leadership=leadership,
        **{f"person_{index}": value.strip()[:300] for index, value in enumerate(people, 1)},
        **{f"facebook_{index}": value.strip()[:300] for index, value in enumerate(facebook_profiles, 1)},
        created_by=user.email,
    ))
    db.commit()
    return RedirectResponse(public_url("/formularios/red-apoyo?guardado=1"), status_code=303)


@app.get("/admin/formularios/datos-estructura", response_class=HTMLResponse)
def structure_records(request: Request, db: Session = Depends(get_db)):
    require_superadmin(request, db)
    rows = db.scalars(select(StructureSubmission).order_by(StructureSubmission.created_at.desc())).all()
    return page(request, "structure_records.html", rows=rows)


@app.get("/admin/formularios/datos-estructura/{record_id}/{side}")
def structure_image(record_id: int, side: str, request: Request, db: Session = Depends(get_db)):
    require_superadmin(request, db)
    row = db.get(StructureSubmission, record_id)
    if not row or side not in {"frente", "reverso"}:
        raise HTTPException(404, "Imagen no encontrada")
    content = row.front_image if side == "frente" else row.back_image
    media_type = row.front_mime if side == "frente" else row.back_mime
    extension = {"image/png": "png", "image/webp": "webp"}.get(media_type, "jpg")
    return Response(
        content=content,
        media_type=media_type,
        headers={"Content-Disposition": f'attachment; filename="estructura-{record_id}-{side}.{extension}"'},
    )


@app.get("/admin/formularios/datos-estructura.xlsx")
def export_structure(request: Request, db: Session = Depends(get_db)):
    require_superadmin(request, db)
    rows = db.scalars(select(StructureSubmission).order_by(StructureSubmission.created_at)).all()
    wb = Workbook()
    ws = wb.active
    ws.title = "Datos de estructura"
    ws.append(["ID", "Correo", "Nombre completo", "Fecha de nacimiento", "Sección electoral",
               "Número telefónico", "Perfil de Facebook", "Carrera o nivel de estudios", "Edad",
               "Dependencia y puesto", "Frente de INE", "Reverso de INE", "Capturó", "Fecha de captura"])
    for row in rows:
        base = str(request.base_url).rstrip("/")
        ws.append([row.id, row.email, row.full_name, row.birth_date, row.electoral_section, row.phone,
                   row.facebook_profile, row.education, row.age, row.workplace,
                   f"{base}/admin/formularios/datos-estructura/{row.id}/frente",
                   f"{base}/admin/formularios/datos-estructura/{row.id}/reverso",
                   row.created_by, row.created_at.replace(tzinfo=None)])
    return workbook_response(wb, "datos_estructura.xlsx")


@app.get("/admin/formularios/red-apoyo", response_class=HTMLResponse)
def support_network_records(request: Request, db: Session = Depends(get_db)):
    require_superadmin(request, db)
    rows = db.scalars(select(SupportNetworkSubmission).order_by(SupportNetworkSubmission.created_at.desc())).all()
    return page(request, "support_network_records.html", rows=rows)


@app.get("/admin/formularios/red-apoyo.xlsx")
def export_support_network(request: Request, db: Session = Depends(get_db)):
    require_superadmin(request, db)
    rows = db.scalars(select(SupportNetworkSubmission).order_by(SupportNetworkSubmission.created_at)).all()
    wb = Workbook()
    ws = wb.active
    ws.title = "Red de apoyo 1x10"
    person_headers = []
    for number in range(1, 11):
        person_headers.extend([f"Nombre Persona {number}", f"Facebook Persona {number}"])
    ws.append(["ID", "Liderazgo", *person_headers, "Capturó", "Fecha de captura"])
    for row in rows:
        person_values = []
        for number in range(1, 11):
            person_values.extend([getattr(row, f"person_{number}"), getattr(row, f"facebook_{number}")])
        ws.append([row.id, row.leadership, *person_values, row.created_by, row.created_at.replace(tzinfo=None)])
    return workbook_response(wb, "red_apoyo_1x10.xlsx")


def workbook_response(wb: Workbook, filename: str) -> StreamingResponse:
    ws = wb.active
    ws.freeze_panes = "A2"
    for column in ws.columns:
        ws.column_dimensions[column[0].column_letter].width = min(
            max(len(str(cell.value or "")) for cell in column) + 2, 55
        )
    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"', "Cache-Control": "no-store"},
    )


@app.post("/ocr", response_class=HTMLResponse)
async def ocr(
    request: Request,
    front: UploadFile = File(...),
    back: UploadFile | None = File(None),
    csrf: str = Form(...),
    db: Session = Depends(get_db),
):
    require_form_access(request, db, FORM_OCR)
    require_csrf(request, csrf)
    merged = {
        "name": "",
        "given_names": "",
        "paternal_surname": "",
        "maternal_surname": "",
        "address": "",
        "municipality": "",
        "curp": "",
        "cic": "",
        "ocr_code": "",
        "phone": "",
        "leader": "",
    }
    back_requested = bool(back and back.filename)
    try:
        uploads = [("front", front)] + ([("back", back)] if back_requested else [])
        for side, upload in uploads:
            image_data = await read_upload(upload)
            fields, raw = extract_image(image_data, side=side)
            logger.info(
                "ocr_metrics side=%s bytes=%s raw_chars=%s fields=%s",
                side,
                len(image_data),
                len(raw),
                sum(bool(value) for value in fields.values()),
            )
            for key, value in fields.items():
                if key in merged and value and not merged[key]:
                    merged[key] = value
    except HTTPException:
        raise
    except Exception:
        logger.exception("ocr_failed_without_image_content")
        merged.update(given_names="", paternal_surname="", maternal_surname="", municipality="")
        return review_page(
            request,
            db,
            merged,
            error="No se pudo leer la imagen. Capture los datos manualmente.",
        )

    given_names = merged["given_names"]
    paternal = merged["paternal_surname"]
    maternal = merged["maternal_surname"]
    if not any((given_names, paternal, maternal)):
        given_names, paternal, maternal = split_ine_name(merged["name"])
    merged.update(
        given_names=given_names,
        paternal_surname=paternal,
        maternal_surname=maternal,
        municipality=(
            canonical_municipality(merged["municipality"])
            or infer_municipality(merged["address"])
        ),
        support_type_id="",
    )
    if not any((merged["name"], merged["address"], merged["curp"])):
        return review_page(
            request,
            db,
            merged,
            error=(
                "No se detectó texto de la INE. Vuelva a tomarla de cerca, con buena luz, "
                "sin reflejos y con la credencial completa dentro del cuadro."
            ),
        )
    if back_requested and not (merged["cic"] and merged["ocr_code"]):
        return review_page(
            request,
            db,
            merged,
            error=(
                "No se pudieron leer completos el CIC y el OCR del reverso. "
                "Puede corregirlos manualmente o volver a tomar la fotografía."
            ),
        )
    return review_page(request, db, merged)


@app.post("/registros", response_class=HTMLResponse)
def save(
    request: Request,
    given_names: str = Form(""),
    paternal_surname: str = Form(""),
    maternal_surname: str = Form(""),
    municipality: str = Form(""),
    support_type_id: str = Form(""),
    curp: str = Form(""),
    cic: str = Form(""),
    ocr_code: str = Form(""),
    phone: str = Form(""),
    leader: str = Form(""),
    # Previous clients may still submit these fields during a rolling update.
    name: str = Form(""),
    address: str = Form(""),
    csrf: str = Form(...),
    db: Session = Depends(get_db),
):
    user = require_form_access(request, db, FORM_OCR)
    require_csrf(request, csrf)

    if not given_names.strip() and name.strip():
        given_names, paternal_surname, maternal_surname = split_ine_name(name)
    given_names = " ".join(given_names.strip().upper().split())
    paternal_surname = " ".join(paternal_surname.strip().upper().split())
    maternal_surname = " ".join(maternal_surname.strip().upper().split())
    curp = curp.strip().upper()
    cic = re.sub(r"\D", "", cic)
    ocr_code = re.sub(r"\D", "", ocr_code)
    selected_leader = canonical_leader(leader)
    phone = re.sub(r"\D", "", phone)

    selected_municipality = canonical_municipality(municipality)
    try:
        selected_support_id = int(support_type_id)
    except (TypeError, ValueError):
        selected_support_id = 0
    support = db.scalar(
        select(SupportType).where(
            SupportType.id == selected_support_id,
            SupportType.active == true(),
        )
    )

    data = {
        "given_names": given_names,
        "paternal_surname": paternal_surname,
        "maternal_surname": maternal_surname,
        "municipality": selected_municipality or municipality,
        "support_type_id": support_type_id,
        "curp": curp,
        "cic": cic,
        "ocr_code": ocr_code,
        "phone": phone,
        "leader": selected_leader or leader,
        "name": name,
        "address": address,
    }
    if not re.fullmatch(r"[A-Z]{4}\d{6}[HM][A-Z]{5}[A-Z0-9]\d", curp):
        return review_page(request, db, data, error="Capture una CURP válida de 18 caracteres.")
    if cic and not re.fullmatch(r"\d{8,15}", cic):
        return review_page(request, db, data, error="Revise el CIC del reverso; debe tener entre 8 y 15 dígitos.")
    if ocr_code and not re.fullmatch(r"\d{12,14}", ocr_code):
        return review_page(request, db, data, error="Revise el OCR del reverso; debe tener entre 12 y 14 dígitos.")
    if phone and len(phone) != 10:
        return review_page(request, db, data, error="El teléfono debe tener 10 dígitos.")
    if not selected_leader:
        return review_page(request, db, data, error="Seleccione un líder del catálogo.")
    if not given_names or not paternal_surname:
        return review_page(request, db, data, error="Capture los nombres y el apellido paterno.")
    if not selected_municipality:
        return review_page(request, db, data, error="Seleccione un municipio válido.")
    if not support:
        return review_page(request, db, data, error="Seleccione un tipo de apoyo activo.")

    conditions = []
    if curp:
        conditions.append(Person.curp == curp)
    if phone:
        conditions.append(Person.phone == phone)
    if cic:
        conditions.append(Person.cic == cic)
    if ocr_code:
        conditions.append(Person.ocr_code == ocr_code)
    duplicates = list(db.scalars(select(Person).where(or_(*conditions))).unique().all()) if conditions else []
    if duplicates:
        visible_duplicates = duplicates if user.is_superadmin else []
        return review_page(
            request,
            db,
            data,
            duplicates=visible_duplicates,
            error=(
                "Posible duplicado: revise las coincidencias antes de continuar."
                if user.is_superadmin
                else "Ya existe una posible coincidencia. Solicite la revisión de un superadministrador."
            ),
        )

    full_name = " ".join(filter(None, (given_names, paternal_surname, maternal_surname)))
    person = Person(
        name=full_name[:180],
        address="",
        given_names=given_names[:120],
        paternal_surname=paternal_surname[:80],
        maternal_surname=maternal_surname[:80],
        municipality=selected_municipality[:120],
        support_type_id=support.id,
        curp=curp[:18],
        cic=cic[:20],
        ocr_code=ocr_code[:20],
        phone=phone[:15],
        leader=selected_leader[:180],
        created_by=user.email,
    )
    db.add(person)
    db.commit()
    return RedirectResponse(public_url("/captura?guardado=1"), status_code=303)


@app.post("/registros/{person_id}/eliminar")
def delete_record(
    person_id: int,
    request: Request,
    csrf: str = Form(...),
    db: Session = Depends(get_db),
):
    require_superadmin(request, db)
    require_csrf(request, csrf)
    person = db.get(Person, person_id)
    if not person:
        raise HTTPException(404, "Registro no encontrado")
    db.delete(person)
    db.commit()
    return RedirectResponse(public_url("/beneficiarios?eliminado=1"), status_code=303)


@app.get("/exportar.xlsx")
def export(request: Request, db: Session = Depends(get_db)):
    require_superadmin(request, db)
    rows = db.scalars(select(Person).order_by(Person.created_at)).unique().all()
    wb = Workbook()
    ws = wb.active
    ws.title = "Registros"
    ws.append([
        "ID",
        "Nombre(s)",
        "Apellido paterno",
        "Apellido materno",
        "CURP",
        "CIC",
        "OCR",
        "Municipio",
        "Código de apoyo",
        "Tipo de apoyo",
        "Número de teléfono",
        "Líder",
        "Capturó",
        "Fecha de captura",
    ])
    for row in rows:
        ws.append([
            row.id,
            row.given_names,
            row.paternal_surname,
            row.maternal_surname,
            row.curp,
            row.cic,
            row.ocr_code,
            row.municipality,
            row.support_type.code if row.support_type else "",
            row.support_type.name if row.support_type else "",
            row.phone,
            row.leader,
            row.created_by,
            row.created_at.replace(tzinfo=None),
        ])
    ws.freeze_panes = "A2"
    for column in ws.columns:
        ws.column_dimensions[column[0].column_letter].width = min(
            max(len(str(cell.value or "")) for cell in column) + 2,
            55,
        )
    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    headers = {
        "Content-Disposition": 'attachment; filename="registros_apoyos.xlsx"',
        "Cache-Control": "no-store",
    }
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=headers,
    )


def users_page(request: Request, db: Session, *, error: str | None = None):
    users = db.scalars(select(AppUser).order_by(AppUser.email)).all()
    active_admins = db.scalar(select(func.count(AppUser.id)).where(
        AppUser.role == ROLE_SUPERADMIN,
        AppUser.active == true(),
    )) or 0
    return page(
        request,
        "users.html",
        users=users,
        roles=(ROLE_SUPERADMIN, ROLE_CAPTURISTA),
        forms=FORM_DEFINITIONS,
        active_admins=active_admins,
        error=error,
    )


@app.get("/admin/usuarios", response_class=HTMLResponse)
def manage_users(request: Request, db: Session = Depends(get_db)):
    require_superadmin(request, db)
    return users_page(request, db)


@app.post("/admin/usuarios", response_class=HTMLResponse)
def create_user(
    request: Request,
    email: str = Form(...),
    password: str = Form(...),
    role: str = Form(...),
    forms: list[str] | None = Form(None),
    csrf: str = Form(...),
    db: Session = Depends(get_db),
):
    actor = require_superadmin(request, db)
    require_csrf(request, csrf)
    normalized_email = email.strip().casefold()
    if not re.fullmatch(r"[^\s@]+@[^\s@]+\.[^\s@]+", normalized_email):
        return users_page(request, db, error="Capture un correo válido.")
    if not 12 <= len(password) <= 200:
        return users_page(request, db, error="La contraseña debe tener entre 12 y 200 caracteres.")
    if role not in VALID_ROLES:
        raise HTTPException(400, "Rol no válido")
    selected_forms = normalized_form_codes(forms)
    if not selected_forms:
        return users_page(request, db, error="Seleccione al menos un formulario para la cuenta.")
    if set(forms or []) - VALID_FORM_CODES:
        raise HTTPException(400, "Formulario no válido")
    if db.scalar(select(AppUser.id).where(func.lower(AppUser.email) == normalized_email)):
        return users_page(request, db, error="Ya existe una cuenta con ese correo.")
    db.add(AppUser(
        email=normalized_email,
        password_hash=password_hash.hash(password),
        role=role,
        active=True,
        form_access=",".join(selected_forms),
        created_by=actor.email,
    ))
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        return users_page(request, db, error="Ya existe una cuenta con ese correo.")
    return RedirectResponse(public_url("/admin/usuarios?creado=1"), status_code=303)


@app.post("/admin/usuarios/{user_id}/formularios", response_class=HTMLResponse)
def change_user_forms(
    user_id: int,
    request: Request,
    forms: list[str] | None = Form(None),
    csrf: str = Form(...),
    db: Session = Depends(get_db),
):
    require_superadmin(request, db)
    require_csrf(request, csrf)
    target = db.get(AppUser, user_id)
    if not target:
        raise HTTPException(404, "Usuario no encontrado")
    selected_forms = normalized_form_codes(forms)
    if not selected_forms:
        return users_page(request, db, error="Seleccione al menos un formulario para la cuenta.")
    if set(forms or []) - VALID_FORM_CODES:
        raise HTTPException(400, "Formulario no válido")
    target.form_access = ",".join(selected_forms)
    db.commit()
    return RedirectResponse(public_url("/admin/usuarios?actualizado=1"), status_code=303)


def ensure_admin_change_is_safe(db: Session, actor: AppUser, target: AppUser, *, removes_admin: bool) -> None:
    if target.id == actor.id and removes_admin:
        raise HTTPException(400, "No puede quitarse su propio acceso de superadministrador")
    if removes_admin and target.is_superadmin and target.active:
        active_admins = db.scalar(select(func.count(AppUser.id)).where(
            AppUser.role == ROLE_SUPERADMIN,
            AppUser.active == true(),
        )) or 0
        if active_admins <= 1:
            raise HTTPException(400, "Debe permanecer al menos un superadministrador activo")


@app.post("/admin/usuarios/{user_id}/rol")
def change_user_role(
    user_id: int,
    request: Request,
    role: str = Form(...),
    csrf: str = Form(...),
    db: Session = Depends(get_db),
):
    actor = require_superadmin(request, db)
    require_csrf(request, csrf)
    if role not in VALID_ROLES:
        raise HTTPException(400, "Rol no válido")
    target = db.get(AppUser, user_id)
    if not target:
        raise HTTPException(404, "Usuario no encontrado")
    ensure_admin_change_is_safe(
        db,
        actor,
        target,
        removes_admin=target.is_superadmin and role != ROLE_SUPERADMIN,
    )
    target.role = role
    db.commit()
    return RedirectResponse(public_url("/admin/usuarios?actualizado=1"), status_code=303)


@app.post("/admin/usuarios/{user_id}/estado")
def change_user_state(
    user_id: int,
    request: Request,
    active: str = Form(...),
    csrf: str = Form(...),
    db: Session = Depends(get_db),
):
    actor = require_superadmin(request, db)
    require_csrf(request, csrf)
    target = db.get(AppUser, user_id)
    if not target:
        raise HTTPException(404, "Usuario no encontrado")
    new_active = active == "1"
    ensure_admin_change_is_safe(
        db,
        actor,
        target,
        removes_admin=target.active and not new_active,
    )
    target.active = new_active
    db.commit()
    return RedirectResponse(public_url("/admin/usuarios?actualizado=1"), status_code=303)


def support_types_page(request: Request, db: Session, *, error: str | None = None):
    items = db.scalars(
        select(SupportType).order_by(SupportType.sort_order, SupportType.name)
    ).all()
    return page(request, "support_types.html", support_types=items, error=error)


@app.get("/admin/tipos-apoyo", response_class=HTMLResponse)
def manage_support_types(request: Request, db: Session = Depends(get_db)):
    require_superadmin(request, db)
    return support_types_page(request, db)


@app.post("/admin/tipos-apoyo", response_class=HTMLResponse)
def create_support_type(
    request: Request,
    code: str = Form(...),
    name: str = Form(...),
    sort_order: int = Form(0),
    csrf: str = Form(...),
    db: Session = Depends(get_db),
):
    require_superadmin(request, db)
    require_csrf(request, csrf)
    stable_code = re.sub(r"[^A-Z0-9_-]", "", code.strip().upper())[:50]
    display_name = " ".join(name.strip().split())[:180]
    if not stable_code or not display_name:
        return support_types_page(request, db, error="Capture un código y un nombre válidos.")
    if db.scalar(select(SupportType.id).where(SupportType.code == stable_code)):
        return support_types_page(request, db, error="Ese código de apoyo ya existe.")
    db.add(SupportType(code=stable_code, name=display_name, active=True, sort_order=sort_order))
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        return support_types_page(request, db, error="Ese código de apoyo ya existe.")
    return RedirectResponse(public_url("/admin/tipos-apoyo?creado=1"), status_code=303)


@app.post("/admin/tipos-apoyo/{support_type_id}")
def update_support_type(
    support_type_id: int,
    request: Request,
    name: str = Form(...),
    sort_order: int = Form(0),
    active: str | None = Form(None),
    csrf: str = Form(...),
    db: Session = Depends(get_db),
):
    require_superadmin(request, db)
    require_csrf(request, csrf)
    item = db.get(SupportType, support_type_id)
    if not item:
        raise HTTPException(404, "Tipo de apoyo no encontrado")
    display_name = " ".join(name.strip().split())[:180]
    if not display_name:
        return support_types_page(request, db, error="El nombre del apoyo es obligatorio.")
    # The code is intentionally immutable because records and external reports
    # can use it as a stable business identifier.
    item.name = display_name
    item.sort_order = sort_order
    item.active = active == "1"
    db.commit()
    return RedirectResponse(public_url("/admin/tipos-apoyo?actualizado=1"), status_code=303)
