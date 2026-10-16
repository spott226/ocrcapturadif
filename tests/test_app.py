import os
import re
from io import BytesIO
from types import SimpleNamespace

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import inspect, select, text

os.environ["DATABASE_URL"] = "sqlite:///./test.sqlite3"
os.environ["SECRET_KEY"] = "test-secret-key-that-is-long-enough"
os.environ["ADMIN_EMAIL"] = "admin@example.test"
os.environ["ADMIN_PASSWORD"] = "a-secure-test-password"

from app.database import Base, SessionLocal, build_database_url, engine
from app.main import app, migrate_early_v2_database, normalize_public_prefix
from app.models import AppUser, Person, StructureSubmission, SupportNetworkSubmission, SupportType
from tools.importar_excel import import_workbook


@pytest.fixture(autouse=True)
def isolated_database():
    Base.metadata.drop_all(engine)
    Base.metadata.create_all(engine)
    yield
    Base.metadata.drop_all(engine)


def csrf_from(response) -> str:
    match = re.search(r'name="csrf" value="([^"]+)"', response.text)
    assert match, response.text[:500]
    return match.group(1)


def login(client: TestClient, email="admin@example.test", password="a-secure-test-password"):
    login_page = client.get("/login")
    response = client.post(
        "/login",
        data={"email": email, "password": password, "csrf": csrf_from(login_page)},
    )
    assert response.status_code == 200
    return response


def create_support(client: TestClient, code="DESPENSA", name="Apoyo alimentario") -> int:
    home = client.get("/")
    response = client.post(
        "/admin/tipos-apoyo",
        data={
            "code": code,
            "name": name,
            "sort_order": 10,
            "csrf": csrf_from(home),
        },
    )
    assert response.status_code == 200
    with SessionLocal() as db:
        return db.scalar(select(SupportType.id).where(SupportType.code == code))


def valid_record(csrf: str, support_type_id: int, **overrides):
    data = {
        "given_names": "PERSONA",
        "paternal_surname": "FICTICIA",
        "maternal_surname": "PRUEBA",
        "municipality": "Aguascalientes",
        "support_type_id": str(support_type_id),
        "curp": "PULA900101MDFRPN09",
        "cic": "1234567890",
        "ocr_code": "1234567890123",
        "phone": "4491234567",
        "leader": "Nalleli Luévano",
        "csrf": csrf,
    }
    data.update(overrides)
    return data


def test_migration_adds_form_access_when_last_login_already_exists():
    AppUser.__table__.drop(engine)
    with engine.begin() as connection:
        connection.execute(text("""
            CREATE TABLE app_users (
                id INTEGER PRIMARY KEY,
                email VARCHAR(254) NOT NULL,
                password_hash VARCHAR(255) NOT NULL,
                role VARCHAR(32) NOT NULL,
                is_active BOOLEAN NOT NULL,
                created_by VARCHAR(254) NOT NULL,
                created_at DATETIME NOT NULL,
                last_login_at DATETIME NULL
            )
        """))

    migrate_early_v2_database()

    columns = {column["name"] for column in inspect(engine).get_columns("app_users")}
    assert "form_access" in columns


def test_migration_adds_separate_facebook_columns_to_existing_network_table():
    SupportNetworkSubmission.__table__.drop(engine)
    person_columns = ", ".join(f"person_{number} VARCHAR(300) NOT NULL" for number in range(1, 11))
    with engine.begin() as connection:
        connection.execute(text(
            "CREATE TABLE support_network_submissions ("
            "id INTEGER PRIMARY KEY, leadership VARCHAR(180) NOT NULL, "
            f"{person_columns}, created_by VARCHAR(254) NOT NULL, created_at DATETIME NOT NULL)"
        ))

    migrate_early_v2_database()

    columns = {column["name"] for column in inspect(engine).get_columns("support_network_submissions")}
    assert {f"facebook_{number}" for number in range(1, 11)}.issubset(columns)


def test_builds_sql_server_windows_auth_url_without_manual_password_encoding():
    sql_settings = SimpleNamespace(
        database_backend="sqlserver",
        database_url="",
        sqlserver_server=r"SERVIDOR\INSTANCIA",
        sqlserver_database="CapturaApoyosDIF",
        sqlserver_auth="windows",
        sqlserver_username="",
        sqlserver_password="",
        sqlserver_driver="ODBC Driver 18 for SQL Server",
        sqlserver_encrypt=True,
        sqlserver_trust_certificate=False,
    )
    url = build_database_url(sql_settings)
    assert url.drivername == "mssql+pyodbc"
    assert url.host == r"SERVIDOR\INSTANCIA"
    assert url.database == "CapturaApoyosDIF"
    assert url.query["Trusted_Connection"] == "yes"
    assert url.query["Encrypt"] == "yes"
    assert url.query["TrustServerCertificate"] == "no"


def test_builds_sql_server_login_url_without_exposing_password_by_default():
    sql_settings = SimpleNamespace(
        database_backend="sqlserver",
        database_url="",
        sqlserver_server="sql.dif.local",
        sqlserver_database="CapturaApoyosDIF",
        sqlserver_auth="sql",
        sqlserver_username="captura_app",
        sqlserver_password="Prueba! con espacios",
        sqlserver_driver="ODBC Driver 18 for SQL Server",
        sqlserver_encrypt=True,
        sqlserver_trust_certificate=True,
    )
    url = build_database_url(sql_settings)
    assert "Prueba" not in str(url)
    assert url.password == "Prueba! con espacios"
    assert "Trusted_Connection" not in url.query


def test_health_and_home_authentication():
    with TestClient(app) as client:
        health = client.get("/salud")
        assert health.status_code == 200
        assert health.json() == {"status": "ok"}
        response = client.get("/", follow_redirects=False)
        assert response.status_code == 303
        assert response.headers["location"] == "/login"


def test_public_prefix_normalization_for_iis_subpath():
    assert normalize_public_prefix("") == ""
    assert normalize_public_prefix("/") == ""
    assert normalize_public_prefix("CapturaApoyosDIF") == "/CapturaApoyosDIF"
    assert normalize_public_prefix("/CapturaApoyosDIF/") == "/CapturaApoyosDIF"


def test_static_files_are_served_from_the_internal_iis_application_root():
    # IIS strips /CapturaApoyosDIF before forwarding the request. The external
    # prefix belongs in generated links, not in the ASGI root_path.
    assert app.root_path == ""
    with TestClient(app) as client:
        response = client.get("/static/style.css")
        assert response.status_code == 200
        assert response.headers["content-type"].startswith("text/css")


def test_superadmin_can_save_export_and_delete_fictitious_record():
    with TestClient(app) as client:
        login(client)
        support_id = create_support(client)
        saved = client.post(
            "/registros",
            data=valid_record(csrf_from(client.get("/")), support_id),
        )
        assert saved.status_code == 200
        records = client.get("/beneficiarios")
        assert "PERSONA FICTICIA PRUEBA" in records.text
        assert "Aguascalientes" in records.text
        assert "Apoyo alimentario" in records.text

        with SessionLocal() as db:
            person = db.scalar(select(Person).where(Person.curp == "PULA900101MDFRPN09"))
            assert person is not None
            assert person.name == "PERSONA FICTICIA PRUEBA"
            assert person.address == ""
            assert person.given_names == "PERSONA"
            assert person.paternal_surname == "FICTICIA"
            assert person.maternal_surname == "PRUEBA"
            assert person.municipality == "Aguascalientes"
            assert person.support_type_id == support_id
            assert person.cic == "1234567890"
            assert person.ocr_code == "1234567890123"
            assert person.leader == "Nalleli Luévano"
            person_id = person.id

        exported = client.get("/exportar.xlsx")
        assert exported.status_code == 200
        workbook = load_workbook(BytesIO(exported.content), read_only=True)
        rows = list(workbook.active.iter_rows(values_only=True))
        assert rows[0][:10] == (
            "ID", "Nombre(s)", "Apellido paterno", "Apellido materno",
            "CURP", "CIC", "OCR", "Municipio",
            "Código de apoyo", "Tipo de apoyo",
        )
        assert rows[1][1:10] == (
            "PERSONA", "FICTICIA", "PRUEBA",
            "PULA900101MDFRPN09", "1234567890", "1234567890123",
            "Aguascalientes", "DESPENSA", "Apoyo alimentario",
        )
        assert exported.headers["content-disposition"] == 'attachment; filename="registros_apoyos.xlsx"'
        assert workbook.active.title == "Registros"
        workbook.close()

        deleted = client.post(
            f"/registros/{person_id}/eliminar",
            data={"csrf": csrf_from(client.get("/"))},
            follow_redirects=False,
        )
        assert deleted.status_code == 303
        with SessionLocal() as db:
            assert db.get(Person, person_id) is None


def test_rejects_invalid_phone_length():
    with TestClient(app) as client:
        login(client)
        support_id = create_support(client)
        response = client.post(
            "/registros",
            data=valid_record(csrf_from(client.get("/")), support_id, phone="449123"),
        )
        assert response.status_code == 200
        assert "El teléfono debe tener 10 dígitos" in response.text


def test_leader_must_come_from_unique_catalog():
    with TestClient(app) as client:
        login(client)
        support_id = create_support(client)
        response = client.post(
            "/registros",
            data=valid_record(
                csrf_from(client.get("/")),
                support_id,
                leader="LÍDER QUE NO EXISTE",
            ),
        )

        assert response.status_code == 200
        assert "Seleccione un líder del catálogo" in response.text
        assert response.text.count(">Jorge Fabian Luévano Rodríguez</option>") == 1
        with SessionLocal() as db:
            assert db.scalar(select(Person.id)) is None


def test_front_and_optional_back_are_merged_into_review(monkeypatch):
    readings = {
        "front": {
            "name": "PRUEBA LOPEZ ANA",
            "given_names": "ANA",
            "paternal_surname": "PRUEBA",
            "maternal_surname": "LOPEZ",
            "address": "CALLE FICTICIA 10\nAGUASCALIENTES AGS",
            "municipality": "AGUASCALIENTES",
            "curp": "PULA900101MDFRPN09",
        },
        "back": {
            "cic": "1382528441",
            "ocr_code": "3904033366874",
        },
    }

    def fake_extract(_data, side=None):
        return readings[side], f"LECTURA FICTICIA {side}"

    monkeypatch.setattr("app.main.extract_image", fake_extract)
    with TestClient(app) as client:
        login(client)
        create_support(client)
        home = client.get("/captura")
        assert 'name="front"' in home.text
        assert 'name="back"' in home.text

        response = client.post(
            "/ocr",
            data={"csrf": csrf_from(home)},
            files={
                "front": ("frente-ficticio.jpg", b"frente", "image/jpeg"),
                "back": ("reverso-ficticio.jpg", b"reverso", "image/jpeg"),
            },
        )

        assert response.status_code == 200
        assert 'name="cic" value="1382528441"' in response.text
        assert 'name="ocr_code" value="3904033366874"' in response.text
        assert 'name="support_type_id" class="scroll-select" size="6"' in response.text
        assert "Deslice dentro de la lista para ver todos los apoyos" in response.text


def test_reverse_codes_are_normalized_before_saving():
    with TestClient(app) as client:
        login(client)
        support_id = create_support(client)
        response = client.post(
            "/registros",
            data=valid_record(
                csrf_from(client.get("/")),
                support_id,
                cic="123 456-7890",
                ocr_code="1234 5678-90123",
            ),
        )

        assert response.status_code == 200
        with SessionLocal() as db:
            person = db.scalar(select(Person).where(Person.curp == "PULA900101MDFRPN09"))
            assert person.cic == "1234567890"
            assert person.ocr_code == "1234567890123"


def test_records_are_paginated_ten_per_page():
    with TestClient(app) as client:
        login(client)
        with SessionLocal() as db:
            db.add_all([
                Person(
                    name=f"PERSONA FICTICIA {number:02d}",
                    given_names="PERSONA",
                    paternal_surname=f"FICTICIA {number:02d}",
                    municipality="Aguascalientes",
                    curp=f"FICTICIO{number:010d}"[:18],
                    created_by="admin@example.test",
                )
                for number in range(1, 12)
            ])
            db.commit()

        first_page = client.get("/beneficiarios")
        assert "11 registros en total" in first_page.text
        assert first_page.text.count('class="record-card"') == 10
        second_page = client.get("/beneficiarios?pagina=2")
        assert "2 / 2" in second_page.text
        assert second_page.text.count('class="record-card"') == 1


def test_capturista_can_capture_but_cannot_view_export_delete_or_administer():
    with TestClient(app) as admin:
        login(admin)
        support_id = create_support(admin)
        response = admin.post(
            "/admin/usuarios",
            data={
                "email": "captura@example.test",
                "password": "captura-segura-123",
                "role": "capturista",
                "forms": "ocr",
                "csrf": csrf_from(admin.get("/")),
            },
        )
        assert response.status_code == 200

        with TestClient(app) as capturista:
            home = login(capturista, "captura@example.test", "captura-segura-123")
            assert "Registros</h2>" not in home.text
            assert "Exportar Excel" not in home.text
            assert capturista.get("/exportar.xlsx").status_code == 403
            assert capturista.get("/admin/usuarios").status_code == 403
            assert capturista.get("/admin/tipos-apoyo").status_code == 403

            saved = capturista.post(
                "/registros",
                data=valid_record(csrf_from(capturista.get("/")), support_id),
            )
            assert saved.status_code == 200
            assert "Registro guardado correctamente" in saved.text
            assert "PERSONA FICTICIA PRUEBA" not in saved.text

            with SessionLocal() as db:
                person_id = db.scalar(select(Person.id))
            forbidden = capturista.post(
                f"/registros/{person_id}/eliminar",
                data={"csrf": csrf_from(capturista.get("/"))},
                follow_redirects=False,
            )
            assert forbidden.status_code == 403


def test_form_permissions_limit_each_capturista_and_update_immediately():
    with TestClient(app) as admin:
        login(admin)
        support_id = create_support(admin)
        created = admin.post(
            "/admin/usuarios",
            data={
                "email": "formularios@example.test",
                "password": "formularios-seguros-123",
                "role": "capturista",
                "forms": "google_1",
                "csrf": csrf_from(admin.get("/")),
            },
        )
        assert created.status_code == 200
        with SessionLocal() as db:
            user = db.scalar(select(AppUser).where(AppUser.email == "formularios@example.test"))
            assert user.form_access == "google_1"
            user_id = user.id

        with TestClient(app) as capturista:
            home = login(capturista, "formularios@example.test", "formularios-seguros-123")
            assert "Datos de estructura" in home.text
            assert "Red De Apoyo 1x10" not in home.text
            assert 'name="front"' not in home.text

            forbidden_ocr = capturista.post(
                "/ocr",
                data={"csrf": csrf_from(home)},
                files={"front": ("frente.jpg", b"frente", "image/jpeg")},
            )
            assert forbidden_ocr.status_code == 403
            forbidden_save = capturista.post(
                "/registros",
                data=valid_record(csrf_from(capturista.get("/")), support_id),
            )
            assert forbidden_save.status_code == 403

            updated = admin.post(
                f"/admin/usuarios/{user_id}/formularios",
                data={
                    "forms": ["ocr", "google_1"],
                    "csrf": csrf_from(admin.get("/")),
                },
            )
            assert updated.status_code == 200
            refreshed = capturista.get("/captura")
            assert 'name="front"' in refreshed.text
            assert "Datos de estructura" in capturista.get("/").text


def test_internal_structure_form_stores_both_ine_images_and_exports():
    with TestClient(app) as client:
        login(client)
        form = client.get("/formularios/datos-estructura")
        assert "← Volver al inicio" in form.text
        saved = client.post(
            "/formularios/datos-estructura",
            data={
                "email": "persona@example.test",
                "full_name": "PERSONA DE PRUEBA",
                "birth_date": "1990-01-01",
                "electoral_section": "0028",
                "phone": "4491234567",
                "facebook_profile": "facebook.com/persona",
                "education": "Licenciatura",
                "age": "36",
                "workplace": "Dependencia / Puesto",
                "csrf": csrf_from(form),
            },
            files={
                "front": ("frente.jpg", b"imagen-frente", "image/jpeg"),
                "back": ("reverso.png", b"imagen-reverso", "image/png"),
            },
        )
        assert saved.status_code == 200
        with SessionLocal() as db:
            row = db.scalar(select(StructureSubmission))
            assert row.front_image == b"imagen-frente"
            assert row.back_image == b"imagen-reverso"
            row_id = row.id
        assert client.get(f"/admin/formularios/datos-estructura/{row_id}/frente").content == b"imagen-frente"
        assert client.get("/admin/formularios/datos-estructura.xlsx").status_code == 200


def test_internal_support_network_uses_its_own_leadership_catalog_and_exports():
    with TestClient(app) as client:
        login(client)
        form = client.get("/formularios/red-apoyo")
        assert "Arath Gomez" in form.text
        assert "← Volver al inicio" in form.text
        assert 'name="leadership" required' in form.text
        assert 'size="7"' not in form.text
        data = {"leadership": "Arath Gomez", "csrf": csrf_from(form)}
        data.update({f"person_{number}": f"Persona {number}" for number in range(1, 11)})
        data.update({f"facebook_{number}": f"facebook.com/persona{number}" for number in range(1, 11)})
        saved = client.post("/formularios/red-apoyo", data=data)
        assert saved.status_code == 200
        with SessionLocal() as db:
            row = db.scalar(select(SupportNetworkSubmission))
            assert row.leadership == "Arath Gomez"
            assert row.person_10 == "Persona 10"
            assert row.facebook_10 == "facebook.com/persona10"
        assert client.get("/admin/formularios/red-apoyo.xlsx").status_code == 200


def test_superadmin_cannot_disable_or_demote_self():
    with TestClient(app) as client:
        login(client)
        with SessionLocal() as db:
            admin_id = db.scalar(select(AppUser.id).where(AppUser.email == "admin@example.test"))
        csrf = csrf_from(client.get("/"))
        assert client.post(
            f"/admin/usuarios/{admin_id}/estado",
            data={"active": "0", "csrf": csrf},
        ).status_code == 400
        assert client.post(
            f"/admin/usuarios/{admin_id}/rol",
            data={"role": "capturista", "csrf": csrf},
        ).status_code == 400


def test_user_can_change_password_and_old_password_stops_working():
    new_password = "nueva-contrasena-segura-456"
    with TestClient(app) as client:
        login(client)
        form = client.get("/cuenta/contrasena")
        response = client.post(
            "/cuenta/contrasena",
            data={
                "current_password": "a-secure-test-password",
                "new_password": new_password,
                "confirm_password": new_password,
                "csrf": csrf_from(form),
            },
        )
        assert response.status_code == 200
        assert "Contraseña actualizada" in response.text

        old_login = client.post(
            "/login",
            data={
                "email": "admin@example.test",
                "password": "a-secure-test-password",
                "csrf": csrf_from(client.get("/login")),
            },
        )
        assert "Correo o contraseña incorrectos" in old_login.text

        new_login = login(client, password=new_password)
        assert "Superadministrador" in new_login.text


def test_excel_import_supports_legacy_and_v2_without_writing_in_simulation(tmp_path):
    with TestClient(app):
        with SessionLocal() as db:
            db.add(SupportType(code="DESPENSA", name="Apoyo alimentario", active=True))
            db.commit()

        legacy = Workbook()
        sheet = legacy.active
        sheet.append(["Nombre completo", "CURP", "Dirección", "Número de teléfono"])
        sheet.append([
            "PERSONA LEGADA FICTICIA", "LEGA900101MDFRPN09",
            "DOMICILIO FICTICIO", "4491111111",
        ])
        legacy_path = tmp_path / "legacy.xlsx"
        legacy.save(legacy_path)
        assert import_workbook(legacy_path, simulate=True) == (1, 0, 0)
        with SessionLocal() as db:
            assert db.scalar(select(Person.id)) is None
        assert import_workbook(legacy_path) == (1, 0, 0)

        modern = Workbook()
        sheet = modern.active
        sheet.append([
            "Nombre(s)", "Apellido paterno", "Apellido materno", "CURP",
            "Municipio", "Código de apoyo", "Número de teléfono",
        ])
        sheet.append([
            "OTRA", "PERSONA", "FICTICIA", "OTRA900101MDFRPN09",
            "Aguascalientes", "DESPENSA", "4492222222",
        ])
        sheet.append([
            "TIPO", "DESCONOCIDO", "FICTICIO", "TIPO900101MDFRPN09",
            "Aguascalientes", "NO-EXISTE", "4493333333",
        ])
        modern_path = tmp_path / "v2.xlsx"
        modern.save(modern_path)
        assert import_workbook(modern_path) == (1, 0, 1)

        with SessionLocal() as db:
            rows = list(db.scalars(select(Person).order_by(Person.id)).all())
            assert len(rows) == 2
            assert rows[0].name == "PERSONA LEGADA FICTICIA"
            assert rows[0].address == "DOMICILIO FICTICIO"
            assert rows[0].support_type_id is None
            assert rows[1].name == "OTRA PERSONA FICTICIA"
            assert rows[1].address == ""
            assert rows[1].support_type.code == "DESPENSA"
